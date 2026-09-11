"""生成「数据报表」类语料示例（xlsx）。

语料文件已随仓库提供；本脚本用于按需重新生成，避免语料目录混入代码。
用法：python scripts/make_sample_reports.py
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

OUT_DIR = Path(__file__).resolve().parents[1] / "data" / "corpus" / "数据报表"

THIN = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)
HEADER_FONT = Font(bold=True, size=11, color="FFFFFF")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATA_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)
MONEY_ALIGN = Alignment(horizontal="right", vertical="center")


def _write_sheet(ws, title, headers, data, header_fill, widths, money_cols=()):
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + col)].width = width

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = Font(bold=True, size=14)
    cell.alignment = Alignment(horizontal="center", vertical="center")

    for col, header in enumerate(headers, 1):
        c = ws.cell(row=3, column=col, value=header)
        c.font = HEADER_FONT
        c.alignment = HEADER_ALIGN
        c.fill = PatternFill(start_color=header_fill, end_color=header_fill, fill_type="solid")
        c.border = THIN

    for row_idx, row in enumerate(data, 4):
        for col_idx, value in enumerate(row, 1):
            c = ws.cell(row=row_idx, column=col_idx, value=value)
            if col_idx in money_cols:
                c.alignment = MONEY_ALIGN
                c.number_format = "#,##0"
            else:
                c.alignment = DATA_ALIGN
            c.border = THIN


def make_travel_standards():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "差旅费标准"
    headers = ["职级", "交通工具", "住宿标准(元/天)", "餐饮补贴(元/天)", "城市级别"]
    data = [
        ["总监及以上", "头等舱/商务舱、一等座", 800, 100, "一线城市"],
        ["总监及以上", "头等舱/商务舱、一等座", 600, 80, "二线城市"],
        ["总监及以上", "头等舱/商务舱、一等座", 400, 60, "三线及以下"],
        ["部门经理", "经济舱、二等座", 600, 100, "一线城市"],
        ["部门经理", "经济舱、二等座", 450, 80, "二线城市"],
        ["部门经理", "经济舱、二等座", 300, 60, "三线及以下"],
        ["普通员工", "经济舱、二等座", 500, 100, "一线城市"],
        ["普通员工", "经济舱、二等座", 350, 80, "二线城市"],
        ["普通员工", "经济舱、二等座", 250, 60, "三线及以下"],
        ["实习生", "经济舱、二等座", 300, 50, "一线城市"],
    ]
    _write_sheet(ws, "差旅费标准表", headers, data, "4472C4",
                 [15, 20, 18, 18, 15], money_cols=(3, 4))
    path = OUT_DIR / "差旅费标准.xlsx"
    wb.save(path)
    print(f"已生成 {path}")


def make_salary_standards():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "薪酬标准"
    headers = ["职级", "基本工资", "绩效工资", "岗位津贴", "年终奖系数"]
    data = [
        ["P1", 6000, 2000, 500, 1.0],
        ["P2", 8000, 3000, 800, 1.2],
        ["P3", 12000, 4000, 1200, 1.5],
        ["P4", 18000, 6000, 1800, 2.0],
        ["M1", 22000, 8000, 2200, 2.5],
        ["M2", 28000, 10000, 2800, 3.0],
        ["M3", 35000, 12000, 3500, 3.5],
        ["M4", 45000, 15000, 4500, 4.0],
    ]
    _write_sheet(ws, "职级薪酬标准表", headers, data, "70AD47",
                 [12, 15, 15, 15, 15], money_cols=(2, 3, 4))
    path = OUT_DIR / "职级薪酬表.xlsx"
    wb.save(path)
    print(f"已生成 {path}")


if __name__ == "__main__":
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    make_travel_standards()
    make_salary_standards()
    print("完成")
